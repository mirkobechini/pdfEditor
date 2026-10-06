"""Upload/import mixin for PdfService (extracted from pdf_service.py)."""
from app.core.config import settings
from app.core.storage import save_pdf, validate_pdf, get_file_content  # noqa: F401
from app.models.pdf import PdfDocument


class PdfServiceUploadMixin:
    """PDF ingestion: upload and file-to-PDF import."""

    def upload(self, filename: str, content: bytes, user_id: str, upload_source: str = "web") -> PdfDocument:
        """Validate, save to disk, and create DB record."""
        # Validate PDF
        if not validate_pdf(content):
            raise ValueError("Invalid PDF file")

        # Get page count and check encryption with PyMuPDF
        import fitz

        doc = fitz.open(stream=content, filetype="pdf")
        page_count = doc.page_count
        is_encrypted = bool(doc.needs_pass)
        # Extract original PDF creation date from metadata
        meta = doc.metadata
        pdf_creation_date = meta.get("creationDate") if isinstance(meta, dict) else None
        doc.close()

        # Enforce page limit (skip for encrypted — can't verify without password)
        if not is_encrypted and page_count > settings.MAX_PAGE_COUNT:
            raise ValueError(
                f"PDF has {page_count} pages. Maximum allowed is {settings.MAX_PAGE_COUNT}"
            )

        # Save to storage
        file_uuid = save_pdf(content)

        # Create DB record
        pdf = PdfDocument(
            original_filename=filename,
            storage_filename=f"{file_uuid}.pdf",
            file_size=len(content),
            page_count=page_count if not is_encrypted else 0,
            is_password_protected=is_encrypted,
            pdf_creation_date=pdf_creation_date,
            upload_source=upload_source,
            user_id=user_id,
        )
        return self.repo.create(pdf)

    def import_file_to_pdf(self, filename: str, content: bytes, user_id: str) -> PdfDocument:
        """Import a file and convert it to PDF."""
        import fitz

        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

        if ext == "txt":
            text = content.decode("utf-8", errors="replace")
            doc = fitz.open()
            page_idx = doc.insert_page(-1, width=612, height=792)
            page = doc[page_idx]
            page.insert_text((50, 100), text, fontname="helv", fontsize=11)
            pdf_bytes = doc.tobytes()
            doc.close()

        elif ext in ("png", "jpg", "jpeg", "gif", "bmp"):
            # Normalize the image to PNG via Pillow before passing to fitz.
            # PyMuPDF does not reliably open GIF/BMP streams directly with
            # filetype="gif"/"bmp" (raises on some builds), and jpg needs
            # filetype="jpeg". Converting to PNG first makes all image
            # formats work uniformly and robustly.
            from io import BytesIO
            from PIL import Image

            try:
                img = Image.open(BytesIO(content))
                img.load()
            except Exception as e:
                raise ValueError(f"Invalid or unsupported image: {e}")

            png_buf = BytesIO()
            # Preserve alpha channel if present (e.g. transparent PNG/GIF).
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
            else:
                img = img.convert("RGB")
            img.save(png_buf, format="PNG")

            doc = fitz.open(stream=png_buf.getvalue(), filetype="png")
            # convert_to_pdf() is required for image documents; tobytes()/save()
            # only works on PDF documents and raises AssertionError otherwise.
            pdf_bytes = doc.convert_to_pdf()
            doc.close()

        elif ext == "docx":
            pdf_bytes = self._convert_docx_to_pdf(content)

        else:
            raise ValueError(f"Unsupported import format: {ext}")

        # Validate and save
        if not validate_pdf(pdf_bytes):
            raise ValueError("Conversion produced an invalid PDF")

        file_uuid = save_pdf(pdf_bytes)
        doc_fitz = fitz.open(stream=pdf_bytes, filetype="pdf")
        page_count = doc_fitz.page_count
        doc_fitz.close()

        pdf = PdfDocument(
            original_filename=filename,
            storage_filename=f"{file_uuid}.pdf",
            file_size=len(pdf_bytes),
            page_count=page_count,
            user_id=user_id,
        )
        return self.repo.create(pdf)

    def _convert_docx_to_pdf(self, content: bytes) -> bytes:
        """Convert a DOCX file to PDF bytes using python-docx + reportlab.

        Reads paragraphs (and basic formatting) from the DOCX and renders
        them into a PDF. Simple but functional — complex layouts (tables,
        images, styles) may lose fidelity.
        """
        from io import BytesIO
        from docx import Document
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas

        try:
            doc = Document(BytesIO(content))
        except Exception as e:
            raise ValueError(f"Invalid DOCX file: {e}")

        buffer = BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        margin = 20 * mm
        y = height - margin
        line_height = 14

        def new_page_if_needed():
            nonlocal y
            if y < margin:
                c.showPage()
                y = height - margin

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                y -= line_height
                continue

            # Basic font size from style (default 11pt)
            font_size = 11
            try:
                if para.style and para.style.font and para.style.font.size:
                    font_size = para.style.font.size.pt
            except Exception:
                pass

            c.setFont("Helvetica", font_size)
            # Wrap text to fit the page width
            max_chars = int((width - 2 * margin) / (font_size * 0.5))
            words = text.split()
            line = ""
            for word in words:
                candidate = f"{line} {word}".strip()
                if len(candidate) > max_chars:
                    new_page_if_needed()
                    c.drawString(margin, y, line)
                    y -= line_height
                    line = word
                else:
                    line = candidate
            if line:
                new_page_if_needed()
                c.drawString(margin, y, line)
                y -= line_height

        c.showPage()
        c.save()
        return buffer.getvalue()